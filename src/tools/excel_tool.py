import re
from copy import copy
from pathlib import Path
from typing import List, Optional, Dict
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from src.models import TestCase, ExecutionResult, RunSummary
from src.utils import get_logger


logger = get_logger(__name__)

_ALIASES = {
    'id': ['testcase id', 'test case id', 'tc id', 'id'],
    'title': ['title', 'test case title'],
    'steps': ['test step', 'test steps', 'steps'],
    'step_no': ['step no', 'step number'],
    'action': ['action'],
    'element': ['element description'],
    'data': ['test data'],
    'expected': ['expected result'],
    'execute': ['execute'],
    'result': ['execution result', 'execution status', 'status'],
}

RESULT_COLUMNS = ['Actual Result', 'Screenshot Path', 'Execution Time (ms)',
                  'Element Source', 'Tokens Used', 'Run ID', 'Executed At']

_STATUS_COLORS = {'Pass': 'C6EFCE', 'Fail': 'FFC7CE', 'Blocked': 'FFEB9C', 'Skipped': 'D9D9D9'}


def _norm(value) -> str:
    return re.sub(r'\s+', ' ', str(value or '')).strip().lower()


def _find_sheet(wb):
    if 'Test Cases' in wb.sheetnames:
        return wb['Test Cases']
    for name in wb.sheetnames:
        if name not in ('Run Summary', 'Step Results'):
            return wb[name]
    return wb.active


def _map_columns(ws) -> Dict[str, int]:
    headers = {_norm(c.value): c.column for c in ws[1] if c.value is not None}
    columns = {}
    for key, aliases in _ALIASES.items():
        for alias in aliases:
            if alias in headers:
                columns[key] = headers[alias]
                break
    return columns


def parse_steps(cell_value) -> List[str]:
    """Split a multi-line 'Test Step' cell into individual manual steps, dropping the leading numbering."""
    steps: List[str] = []
    for line in str(cell_value or '').splitlines():
        line = line.strip()
        if not line:
            continue
        match = re.match(r'^\s*\d+\s*[.)\-:]\s*(.*)$', line)
        if match:
            steps.append(match.group(1).strip())
        elif steps:
            steps[-1] += ' ' + line
        else:
            steps.append(line)
    return steps


class ReadExcelTool:
    def __init__(self, file_path: str):
        self.file_path = Path(file_path)

    def read_test_suite(self, application: str, environment: str) -> List[TestCase]:
        if not self.file_path.exists():
            raise FileNotFoundError(f"Test suite file not found: {self.file_path}")

        wb = load_workbook(self.file_path, data_only=True)
        ws = _find_sheet(wb)
        cols = _map_columns(ws)
        if 'id' not in cols or ('steps' not in cols and 'action' not in cols):
            raise ValueError("Test suite needs 'TestCase ID' and 'Test Step' columns (or 'Step No' + 'Action')")

        def cell(row, key):
            idx = cols.get(key)
            return row[idx - 1].value if idx else None

        test_cases: Dict[str, TestCase] = {}
        order: List[str] = []

        for row in ws.iter_rows(min_row=2):
            tc_id = cell(row, 'id')
            if tc_id is None or not str(tc_id).strip():
                continue
            tc_id = str(tc_id).strip()
            execute = str(cell(row, 'execute') or '').strip().upper() == 'Y'

            if 'step_no' in cols:
                # one row per step: group rows by test case id
                action = str(cell(row, 'action') or '').strip()
                element = str(cell(row, 'element') or '').strip()
                step_text = str(cell(row, 'steps') or '').strip() or f"{action} {element}".strip()
                tc = test_cases.get(tc_id)
                if tc is None:
                    tc = TestCase(test_case_id=tc_id, test_case_title=str(cell(row, 'title') or tc_id),
                                  application=application, environment=environment,
                                  row_index=row[0].row, execute=False)
                    test_cases[tc_id] = tc
                    order.append(tc_id)
                tc.execute = tc.execute or execute
                if execute and step_text:
                    tc.source_steps.append(step_text)
                data = cell(row, 'data')
                if data:
                    tc.test_data = f"{tc.test_data}\n{data}" if tc.test_data else str(data)
                expected = cell(row, 'expected')
                if expected:
                    tc.expected_result = str(expected).strip()
            else:
                title = str(cell(row, 'title') or tc_id).strip()
                test_cases[tc_id] = TestCase(
                    test_case_id=tc_id, test_case_title=title,
                    application=application, environment=environment,
                    source_steps=parse_steps(cell(row, 'steps')),
                    test_data=str(cell(row, 'data')).strip() if cell(row, 'data') else None,
                    expected_result=str(cell(row, 'expected') or '').strip(),
                    execute=execute, row_index=row[0].row)
                order.append(tc_id)

        wb.close()
        return [test_cases[t] for t in order]


class WriteExcelTool:
    def __init__(self, file_path: str):
        self.file_path = Path(file_path)

    def write_all(self, execution_results: List[ExecutionResult], run_summary: RunSummary) -> str:
        """Write results into the same workbook. Returns the path actually written."""
        wb = load_workbook(self.file_path)
        ws = _find_sheet(wb)
        cols = _map_columns(ws)

        header_style = ws.cell(row=1, column=1)
        next_col = ws.max_column + 1

        def ensure_column(header: str, existing: Optional[int] = None) -> int:
            nonlocal next_col
            if existing:
                return existing
            for c in ws[1]:
                if _norm(c.value) == _norm(header):
                    return c.column
            col = next_col
            next_col += 1
            target = ws.cell(row=1, column=col, value=header)
            if header_style.has_style:
                target.font = copy(header_style.font)
                target.fill = copy(header_style.fill)
                target.border = copy(header_style.border)
                target.alignment = copy(header_style.alignment)
            ws.column_dimensions[get_column_letter(col)].width = 28 if header == 'Actual Result' else 18
            return col

        status_col = ensure_column('Status', cols.get('result'))
        out_cols = {h: ensure_column(h) for h in RESULT_COLUMNS}

        for result in execution_results:
            if not result.row_index:
                continue
            r = result.row_index
            status_cell = ws.cell(row=r, column=status_col, value=result.status)
            color = _STATUS_COLORS.get(result.status)
            if color:
                status_cell.fill = PatternFill(start_color=color, end_color=color, fill_type='solid')
            sources = ', '.join(f"{k}:{v}" for k, v in result.element_sources.items()) if result.element_sources else ''
            values = {
                'Actual Result': result.actual_result,
                'Screenshot Path': result.screenshot_path,
                'Execution Time (ms)': round(result.total_execution_time_ms, 1),
                'Element Source': sources,
                'Tokens Used': result.total_tokens_used,
                'Run ID': result.run_id,
                'Executed At': result.completed_at.isoformat(timespec='seconds'),
            }
            for header, value in values.items():
                c = ws.cell(row=r, column=out_cols[header], value=value)
                if header == 'Actual Result':
                    c.alignment = Alignment(wrap_text=True, vertical='top')

        self._write_step_sheet(wb, execution_results)
        self._write_summary_sheet(wb, run_summary)

        try:
            wb.save(self.file_path)
            written = str(self.file_path)
        except PermissionError:
            fallback = self.file_path.with_name(f"{self.file_path.stem}_results{self.file_path.suffix}")
            wb.save(fallback)
            written = str(fallback)
            logger.warning(f"{self.file_path.name} is open in another program; results saved to {fallback.name}")
        wb.close()
        return written

    @staticmethod
    def _write_step_sheet(wb, execution_results: List[ExecutionResult]):
        rerun_ids = {r.test_case_id for r in execution_results}
        kept_rows = []
        if 'Step Results' in wb.sheetnames:
            # a partial run (--only) must not erase the step details of tests that were not re-run
            kept_rows = [row for row in wb['Step Results'].iter_rows(min_row=2, values_only=True)
                         if row[0] and row[0] not in rerun_ids]
            del wb['Step Results']
        ws = wb.create_sheet('Step Results')
        headers = ['Test Case ID', 'Step No', 'Action', 'Target', 'Status', 'Actual Result',
                   'Element Source', 'Tokens Used', 'Execution Time (ms)', 'Screenshot Path', 'Executed At']
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        rows = list(kept_rows)
        for result in execution_results:
            for s in result.steps:
                rows.append((result.test_case_id, s.step_no, s.action, s.target, s.status, s.actual_result,
                             s.element_source, s.tokens_used, round(s.execution_time_ms, 1),
                             s.screenshot_path, s.executed_at.isoformat(timespec='seconds')))
        rows.sort(key=lambda r: (str(r[0]), r[1] or 0))
        for row in rows:
            ws.append(list(row))
            color = _STATUS_COLORS.get(row[4])
            if color:
                ws.cell(row=ws.max_row, column=5).fill = PatternFill(start_color=color, end_color=color, fill_type='solid')
        for i, width in enumerate([12, 8, 12, 50, 10, 60, 14, 12, 18, 40, 20], start=1):
            ws.column_dimensions[get_column_letter(i)].width = width

    @staticmethod
    def _write_summary_sheet(wb, s: RunSummary):
        if 'Run Summary' in wb.sheetnames:
            del wb['Run Summary']
        ws = wb.create_sheet('Run Summary')
        rows = [
            ('Run ID', s.run_id), ('Suite Name', s.suite_name), ('Application', s.application),
            ('Environment', s.environment), ('Start Time', s.start_time.isoformat(timespec='seconds')),
            ('End Time', s.end_time.isoformat(timespec='seconds')),
            ('Total Duration (s)', round(s.total_duration_ms / 1000, 2)), ('', ''),
            ('Test Cases', 'Count'), ('Total', s.total_test_cases), ('Passed', s.passed_test_cases),
            ('Failed', s.failed_test_cases), ('Blocked', s.blocked_test_cases),
            ('Skipped', s.skipped_test_cases), ('', ''),
            ('Steps', 'Count'), ('Total', s.total_steps), ('Passed', s.passed_steps),
            ('Failed', s.failed_steps), ('Blocked', s.blocked_steps), ('Skipped', s.skipped_steps), ('', ''),
            ('LLM Usage', 'Value'), ('Total Tokens', s.total_tokens_used), ('LLM Calls', s.total_llm_calls),
            ('Cache Hits', s.cache_hit_count), ('Fallback Hits', s.fallback_hit_count),
            ('AI Identifications', s.ai_identification_count), ('Heals', s.healed_count),
            ('Cache Hit Rate (cache+fallback / lookups)', f"{s.cache_hit_rate:.2%}"),
            ('False Heal Rate', f"{s.false_heal_rate:.2%}"), ('', ''),
            ('Cost', 'Value'), ('Total Cost (USD)', round(s.total_cost_usd, 6)),
            ('Cost per 100 Steps (USD)', round(s.cost_per_100_steps, 6)),
        ]
        for label, value in rows:
            ws.append([label, value])
        for label in ('Test Cases', 'Steps', 'LLM Usage', 'Cost'):
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
                if row[0].value == label:
                    for c in row[:2]:
                        c.font = Font(bold=True, color='FFFFFF')
                        c.fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        ws.column_dimensions['A'].width = 42
        ws.column_dimensions['B'].width = 30
